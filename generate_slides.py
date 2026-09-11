import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(12, 12, 12)       # Obsidian #0c0c0c
    CARD_BG = RGBColor(22, 22, 22)        # Charcoal #161616
    GOLD = RGBColor(212, 175, 55)         # Warm Gold #d4af37
    GOLD_LIGHT = RGBColor(230, 202, 101)  # Light Champagne #e6ca65
    GOLD_MUTED = RGBColor(197, 168, 128)  # Muted Gold #c5a880
    WHITE = RGBColor(245, 245, 245)       # Off-white
    GREY = RGBColor(160, 160, 160)        # Muted text
    BORDER_GOLD = RGBColor(80, 65, 25)

    # Asset paths
    base_dir = os.path.dirname(os.path.abspath(__file__))
    img_dir = os.path.join(base_dir, "assets", "images")
    logo_path = os.path.join(img_dir, "logo.jpg")
    hero_path = os.path.join(img_dir, "hero.jpg")
    noir_path = os.path.join(img_dir, "noir.jpg")
    bloom_path = os.path.join(img_dir, "bloom.jpg")
    aqua_path = os.path.join(img_dir, "aqua.jpg")
    oud_path = os.path.join(img_dir, "oud.jpg")

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background() # no line
        return bg

    def add_header(slide, tag_text, title_text):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = GOLD_MUTED

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Gold accent line under header
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.5), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD
        line.line.fill.background()

    # ========================================================
    # SLIDE 1: Title Slide
    # ========================================================
    blank_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Accent frame
    frame = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
    frame.fill.solid()
    frame.fill.fore_color.rgb = CARD_BG
    frame.line.color.rgb = GOLD
    frame.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(5.4), Inches(1.1), width=Inches(2.5))

    t_box1 = s1.shapes.add_textbox(Inches(1.0), Inches(3.4), Inches(11.333), Inches(1.2))
    tf1 = t_box1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "V É L O R A   P A R F U M S"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_LIGHT
    p1.alignment = PP_ALIGN.CENTER

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Haute Parfumerie E-Commerce Web Application"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = WHITE
    p1_sub.alignment = PP_ALIGN.CENTER

    meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.333), Inches(1.4))
    tf_meta = meta_box.text_frame
    p_meta1 = tf_meta.paragraphs[0]
    p_meta1.text = "Course: DTI 224 — Web Development Final Project"
    p_meta1.font.size = Pt(14)
    p_meta1.font.color.rgb = GOLD_MUTED
    p_meta1.alignment = PP_ALIGN.CENTER

    p_meta2 = tf_meta.add_paragraph()
    p_meta2.text = "Developer: RexX0Dora1005  •  Technology: HTML5, CSS3, JavaScript (DOM/BOM)"
    p_meta2.font.size = Pt(13)
    p_meta2.font.color.rgb = GREY
    p_meta2.alignment = PP_ALIGN.CENTER

    # ========================================================
    # SLIDE 2: Project Vision & Concept
    # ========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "01. Introduction", "Project Vision & Brand Concept")

    # Left Card
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_GOLD
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Why VÉLORA Luxury Perfume?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD

    points1 = [
        ("Haute Parfumerie Identity", "Simulates an authentic French boutique experience inspired by houses like Chanel, Dior Privée, and Creed."),
        ("Comprehensive Flow", "Demonstrates the complete customer journey: discovery, olfactory curation, sizing, cart management, and official receipting."),
        ("Interactive Engagement", "More than a static catalogue: features an interactive Scent Sommelier quiz and bespoke calligraphy gifting.")
    ]
    for title, desc in points1:
        p_t = tf_c1.add_paragraph()
        p_t.text = "✦  " + title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_d = tf_c1.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = GREY

    # Right Card
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_GOLD
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Academic & Technical Goals"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD

    points2 = [
        ("DTI 224 Objectives", "Applies core curriculum: semantic HTML5 tags, CSS3 modern flexbox & grid architecture, and BOM/DOM event handling."),
        ("Zero Dependency Architecture", "Written in pure vanilla JavaScript without bulky frameworks, achieving blazing speed and clean code readability."),
        ("State Management", "Maintains shopping bag and discount state seamlessly using the browser's localStorage API.")
    ]
    for title, desc in points2:
        p_t = tf_c2.add_paragraph()
        p_t.text = "✦  " + title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_d = tf_c2.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = GREY

    # ========================================================
    # SLIDE 3: Brand Identity & Campaign
    # ========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "02. Visual Design", "Brand Identity & Editorial Aesthetics")

    # Left: Hero Banner Image
    if os.path.exists(hero_path):
        s3.shapes.add_picture(hero_path, Inches(0.8), Inches(1.8), width=Inches(6.4))

    # Right: Design System details
    ds = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.8), Inches(5.0), Inches(5.0))
    ds.fill.solid()
    ds.fill.fore_color.rgb = CARD_BG
    ds.line.color.rgb = BORDER_GOLD
    tf_ds = ds.text_frame
    tf_ds.word_wrap = True
    p = tf_ds.paragraphs[0]
    p.text = "Visual System Architecture"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD

    ds_items = [
        ("Palette: Obsidian & Gold", "Primary background #0a0a0a, deep charcoal surfaces, and warm metallic gold #d4af37 accents."),
        ("Typography Pairings", "Cormorant Garamond (Editorial luxury serif for titles) paired with Montserrat (clean legible sans for UI)."),
        ("Glassmorphism & Lighting", "Backdrop blurs (12px), subtle gold glows, smooth cubic-bezier hover transitions, and marble podium depth."),
        ("Atmospheric Sillage", "Product photos designed with mood-specific lighting: smoky noir, soft rose petals, water ripples, and amber haze.")
    ]
    for t, d in ds_items:
        pt = tf_ds.add_paragraph()
        pt.text = "✦  " + t
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pd = tf_ds.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = GREY

    # ========================================================
    # SLIDE 4: The 4 Fragrance Editions
    # ========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "03. Product Catalog", "The Four Olfactory Masterpieces")

    fragrances = [
        ("VÉLORA NOIR", "Dark & Mysterious", "฿1,290", noir_path, "Pepper, Smoked Cedar, Tonka, Amber"),
        ("VÉLORA BLOOM", "Floral & Sweet", "฿1,190", bloom_path, "Lychee, Damask Rose, Jasmine, Sandalwood"),
        ("VÉLORA AQUA", "Fresh & Oceanic", "฿1,090", aqua_path, "Sea Salt, Mandarin, Marine Accord, Driftwood"),
        ("VÉLORA OUD", "Warm & Woody", "฿1,390", oud_path, "Saffron, Cambodian Oud, Rosewood, Vetiver")
    ]

    col_w = Inches(2.75)
    gap = Inches(0.2)
    start_x = Inches(0.8)

    for i, (name, tag, price, img, notes) in enumerate(fragrances):
        x = start_x + i * (col_w + gap)
        # Card background
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), col_w, Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_GOLD

        # Image
        if os.path.exists(img):
            s4.shapes.add_picture(img, x + Inches(0.2), Inches(2.0), width=Inches(2.35))

        # Text below image
        tbox = s4.shapes.add_textbox(x + Inches(0.15), Inches(4.5), col_w - Inches(0.3), Inches(2.3))
        tf = tbox.text_frame
        tf.word_wrap = True
        pn = tf.paragraphs[0]
        pn.text = name
        pn.font.size = Pt(14)
        pn.font.bold = True
        pn.font.color.rgb = GOLD

        pt = tf.add_paragraph()
        pt.text = tag
        pt.font.size = Pt(10)
        pt.font.color.rgb = GREY

        pp = tf.add_paragraph()
        pp.text = f"Price: {price}"
        pp.font.size = Pt(13)
        pp.font.bold = True
        pp.font.color.rgb = WHITE

        pnot = tf.add_paragraph()
        pnot.text = "Notes: " + notes
        pnot.font.size = Pt(9.5)
        pnot.font.color.rgb = GOLD_MUTED

    # ========================================================
    # SLIDE 5: Technical Architecture
    # ========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "04. Engineering", "Front-End Technical Architecture")

    tech_cols = [
        ("Semantic HTML5", [
            "Modern landmark tags (<header>, <section>, <article>, <aside>, <footer>)",
            "Dual modal architecture (Quick View Modal, 2-Column Checkout)",
            "Native form validation attributes (required, pattern, email)",
            "Custom inline SVG assets for PromptPay QR and luxury iconography"
        ]),
        ("Modern CSS3", [
            "Custom Properties (CSS variables) for consistent black & gold theming",
            "CSS Grid & Flexbox for multi-screen responsive design",
            "Hardware-accelerated animations (smooth slide-in cart, modal scale)",
            "Dedicated @media print stylesheet for A4 printable official receipts"
        ]),
        ("JavaScript (DOM/BOM)", [
            "Pure Vanilla JavaScript without external framework overhead",
            "Dynamic catalog rendering from structured product data objects",
            "Event-driven cart state management synchronized with localStorage",
            "Interactive Scent Sommelier quiz recommendation engine"
        ])
    ]

    card_w = Inches(3.7)
    card_gap = Inches(0.3)
    for i, (title, items) in enumerate(tech_cols):
        cx = Inches(0.8) + i * (card_w + card_gap)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.8), card_w, Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_GOLD
        tfc = c.text_frame
        tfc.word_wrap = True
        p = tfc.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = GOLD

        for item in items:
            pi = tfc.add_paragraph()
            pi.text = "•  " + item
            pi.font.size = Pt(12)
            pi.font.color.rgb = WHITE

    # ========================================================
    # SLIDE 6: Interactive Scent Sommelier Quiz
    # ========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "05. Smart Feature", "AI Scent Sommelier Recommendation Engine")

    q_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    q_box.fill.solid()
    q_box.fill.fore_color.rgb = CARD_BG
    q_box.line.color.rgb = BORDER_GOLD
    tf_q = q_box.text_frame
    tf_q.word_wrap = True

    pq = tf_q.paragraphs[0]
    pq.text = "Elevating E-Commerce with Interactive Personalization"
    pq.font.size = Pt(19)
    pq.font.bold = True
    pq.font.color.rgb = GOLD

    q_points = [
        ("The Business Problem:", "Fragrance is an invisible luxury. Customers online cannot smell the perfumes directly, creating hesitation."),
        ("The Technical Solution:", "An intuitive Scent Sommelier widget. Users select their desired emotional atmosphere (Dark, Sweet, Marine, Woody)."),
        ("Algorithm Matching:", "JavaScript evaluates the choice, selects the matching edition object, renders scent notes, and generates an instant 'Add to Bag' button."),
        ("Customer Value:", "Transforms a passive browsing session into an engaging interactive boutique consultation, significantly boosting conversion rates.")
    ]

    for head, body in q_points:
        ph = tf_q.add_paragraph()
        ph.text = "✦  " + head
        ph.font.size = Pt(14)
        ph.font.bold = True
        ph.font.color.rgb = WHITE
        pb = tf_q.add_paragraph()
        pb.text = "    " + body
        pb.font.size = Pt(12.5)
        pb.font.color.rgb = GREY

    # ========================================================
    # SLIDE 7: 2-Column Luxury Checkout
    # ========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "06. Checkout Flow", "Bespoke 2-Column Checkout & Voucher Engine")

    c7_left = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c7_left.fill.solid()
    c7_left.fill.fore_color.rgb = CARD_BG
    c7_left.line.color.rgb = BORDER_GOLD
    tf_c7l = c7_left.text_frame
    tf_c7l.word_wrap = True
    pl = tf_c7l.paragraphs[0]
    pl.text = "Client & Multi-Channel Payment"
    pl.font.size = Pt(17)
    pl.font.bold = True
    pl.font.color.rgb = GOLD

    l_items = [
        ("Bespoke Delivery Choice", "Standard White-Glove (Complimentary) vs. VIP Same-Day Chauffeur (+฿150) with instant price calculation."),
        ("Personalized Calligraphy", "Complimentary gift option toggle: input recipient name and custom handwritten message for recipient."),
        ("PromptPay QR & Slip Upload", "Generates Thai Banking QR graphic, Kasikornbank account details, and simulated slip attachment preview."),
        ("Smart Card Formatting", "Automatically formats 16 digits into spaced 4-digit blocks (#### #### #### ####) with expiry and CVV validation.")
    ]
    for t, d in l_items:
        pt = tf_c7l.add_paragraph()
        pt.text = "✦  " + t
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pd = tf_c7l.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = GREY

    c7_right = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    c7_right.fill.solid()
    c7_right.fill.fore_color.rgb = CARD_BG
    c7_right.line.color.rgb = BORDER_GOLD
    tf_c7r = c7_right.text_frame
    tf_c7r.word_wrap = True
    pr = tf_c7r.paragraphs[0]
    pr.text = "Voucher Engine & Bag Summary"
    pr.font.size = Pt(17)
    pr.font.bold = True
    pr.font.color.rgb = GOLD

    r_items = [
        ("Live Manifest Preview", "Itemized thumbnails, edition names, and flacon sizes displayed directly alongside order totals."),
        ("Interactive Voucher Logic", "Supports promo codes: 'VELORA10' (10% Off), 'VIP300' (฿300 Off), and 'MAISON' (15% Off)."),
        ("Real-time Discount Recalculation", "Subtotal, promotional deduction, delivery surcharge, and final total calculate synchronously upon application."),
        ("Security Assurance", "Displays 256-bit encryption and authentic maison assurance badges.")
    ]
    for t, d in r_items:
        pt = tf_c7r.add_paragraph()
        pt.text = "✦  " + t
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pd = tf_c7r.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = GREY

    # ========================================================
    # SLIDE 8: Deluxe Official Receipt & Print
    # ========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "07. Confirmation & Output", "Deluxe Digital Certificate & Invoice")

    c8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    c8.fill.solid()
    c8.fill.fore_color.rgb = CARD_BG
    c8.line.color.rgb = BORDER_GOLD
    tfc8 = c8.text_frame
    tfc8.word_wrap = True
    p = tfc8.paragraphs[0]
    p.text = "Authentic Parisian Haute Parfumerie Documentation"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = GOLD

    rec_items = [
        ("Unique Order Reference Ribbon", "Generates unique tracking numbers (#VEL-XXXXXX), exact timestamps, and fulfillment progress."),
        ("Itemized Commercial Manifest", "Shows flacon images, Extrait concentration spec, sizes, unit pricing, and line totals."),
        ("Paris Wax Seal Watermark", "Artisanal extraction guarantee badge assuring formulation in Grasse, France."),
        ("Calligraphy Inscription Display", "Renders the customer's personal gift note in elegant italicized typography."),
        ("One-Click Print Receipt (@media print)", "CSS print stylesheet strips dark backgrounds and headers, outputting a clean, professional A4 PDF invoice."),
        ("Client-side Text Export (.txt)", "Utilizes Blob and URL.createObjectURL to instantly download a formatted plaintext order record.")
    ]
    for t, d in rec_items:
        pt = tfc8.add_paragraph()
        pt.text = "✦  " + t
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pd = tfc8.add_paragraph()
        pd.text = "    " + d
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = GREY

    # ========================================================
    # SLIDE 9: Live Deployment & Links
    # ========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "08. Deployment", "Repository & Live Website Links")

    s9_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    s9_box.fill.solid()
    s9_box.fill.fore_color.rgb = CARD_BG
    s9_box.line.color.rgb = BORDER_GOLD
    tfs9 = s9_box.text_frame
    tfs9.word_wrap = True

    p = tfs9.paragraphs[0]
    p.text = "Production-Ready Deployment"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = GOLD

    dep_points = [
        ("GitHub Source Code Repository", "https://github.com/RexX0Dora1005/VELORA-PERFUME", "Complete source tree with commits, standardized asset images, and clean project documentation."),
        ("Live GitHub Pages URL", "https://rexx0dora1005.github.io/VELORA-PERFUME/", "Fully hosted live web app accessible on any mobile device, tablet, or desktop browser."),
        ("Zero Build Step Architecture", "Vanilla HTML5 / CSS3 / JavaScript", "No node_modules, bundlers, or compilation required — runs directly in vanilla browsers.")
    ]
    for label, link, desc in dep_points:
        pl = tfs9.add_paragraph()
        pl.text = "✦  " + label + ":"
        pl.font.size = Pt(14)
        pl.font.bold = True
        pl.font.color.rgb = WHITE
        
        plk = tfs9.add_paragraph()
        plk.text = "    " + link
        plk.font.size = Pt(13)
        plk.font.bold = True
        plk.font.color.rgb = GOLD_LIGHT

        pd = tfs9.add_paragraph()
        pd.text = "    " + desc
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = GREY

    # ========================================================
    # SLIDE 10: Conclusion & Q&A
    # ========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)

    c10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.6), Inches(12.133), Inches(6.3))
    c10.fill.solid()
    c10.fill.fore_color.rgb = CARD_BG
    c10.line.color.rgb = GOLD
    c10.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s10.shapes.add_picture(logo_path, Inches(5.4), Inches(1.2), width=Inches(2.5))

    t10 = s10.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.333), Inches(1.2))
    tf10 = t10.text_frame
    p10 = tf10.paragraphs[0]
    p10.text = "Thank You for Your Attention"
    p10.font.size = Pt(36)
    p10.font.bold = True
    p10.font.color.rgb = GOLD_LIGHT
    p10.alignment = PP_ALIGN.CENTER

    p10_sub = tf10.add_paragraph()
    p10_sub.text = "“Discover Your Signature Scent — VÉLORA Parfums”"
    p10_sub.font.size = Pt(18)
    p10_sub.font.italic = True
    p10_sub.font.color.rgb = WHITE
    p10_sub.alignment = PP_ALIGN.CENTER

    p10_qa = tf10.add_paragraph()
    p10_qa.text = "Questions & Feedback Welcomed"
    p10_qa.font.size = Pt(16)
    p10_qa.font.bold = True
    p10_qa.font.color.rgb = GOLD_MUTED
    p10_qa.alignment = PP_ALIGN.CENTER

    out_path = os.path.join(base_dir, "VELORA_Presentation.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_deck()
